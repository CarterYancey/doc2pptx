"""Document readers for txt, markdown, docx, pdf, html, and spreadsheets."""

from __future__ import annotations

import os
import time
from pathlib import Path

from .logging_config import _preview, logger
from .models import DeckContent, SlideContent


def read_txt(path: str) -> str:
    """Read plain text file."""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def read_markdown(path: str) -> str:
    """Read markdown file (returns raw markdown text)."""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def read_docx(path: str) -> str:
    """Extract text from a .docx file, preserving heading markers."""
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        style_name = (para.style.name or "").lower()
        # Map Word heading styles to markdown-style markers
        if "heading 1" in style_name:
            lines.append(f"# {text}")
        elif "heading 2" in style_name:
            lines.append(f"## {text}")
        elif "heading 3" in style_name:
            lines.append(f"### {text}")
        elif "heading 4" in style_name:
            lines.append(f"#### {text}")
        elif "title" in style_name:
            lines.append(f"# {text}")
        elif "subtitle" in style_name:
            lines.append(f"_subtitle: {text}")
        elif "list" in style_name:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    return "\n".join(lines)


def read_pdf(path: str) -> str:
    """Extract text from a PDF file."""
    try:
        import pdfplumber
        texts = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    texts.append(page_text)
        if texts:
            return "\n\n".join(texts)
    except Exception:
        pass

    # Fallback to pypdf
    from pypdf import PdfReader
    reader = PdfReader(path)
    texts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            texts.append(page_text)
    return "\n\n".join(texts)


def read_html(path: str) -> str:
    """Extract text from an HTML file, mapping headings to markdown markers."""
    from bs4 import BeautifulSoup

    with open(path, "r", encoding="utf-8", errors="replace") as f:
        soup = BeautifulSoup(f.read(), "html.parser")

    lines = []
    for el in soup.find_all(["h1", "h2", "h3", "h4", "p", "li"]):
        text = el.get_text(strip=True)
        if not text:
            continue
        tag = el.name.lower()
        if tag == "h1":
            lines.append(f"# {text}")
        elif tag == "h2":
            lines.append(f"## {text}")
        elif tag == "h3":
            lines.append(f"### {text}")
        elif tag == "h4":
            lines.append(f"#### {text}")
        elif tag == "li":
            lines.append(f"- {text}")
        else:
            lines.append(text)
        lines.append("")  # blank line after each block element
    return "\n".join(lines)


# ─── Spreadsheet reader (returns DeckContent, not raw text) ──

MAX_TABLE_ROWS_PER_SLIDE = 15   # data rows per table slide (excl. header)
MAX_TABLE_COLS = 10             # columns before we truncate
MAX_COL_CHAR_WIDTH = 30         # truncate long cell values


def _sanitize_cell(val) -> str:
    """Convert a cell value to a clean display string."""
    if val is None:
        return ""
    import math
    if isinstance(val, float):
        if math.isnan(val):
            return ""
        # Show ints cleanly, floats with reasonable precision
        if val == int(val):
            return str(int(val))
        return f"{val:.2f}"
    s = str(val).strip()
    if len(s) > MAX_COL_CHAR_WIDTH:
        s = s[: MAX_COL_CHAR_WIDTH - 1] + "…"
    return s


def read_xlsx(path: str, deck_title: str = "", max_rows_per_slide: int = MAX_TABLE_ROWS_PER_SLIDE) -> DeckContent:
    """
    Read an xlsx/xls/csv/tsv file and produce a DeckContent with table slides.

    Each worksheet becomes one or more table slides. If a sheet has more rows
    than max_rows_per_slide, it is split across multiple slides.
    """
    import pandas as pd

    ext = Path(path).suffix.lower()

    # Read all sheets
    if ext in (".csv", ".tsv"):
        sep = "\t" if ext == ".tsv" else ","
        sheets = {"Sheet1": pd.read_csv(path, sep=sep)}
    else:
        sheets = pd.read_excel(path, sheet_name=None)

    file_stem = Path(path).stem
    deck = DeckContent(title=deck_title or file_stem)

    # Title slide
    deck.slides.append(SlideContent(
        title=deck.title,
        slide_type="title",
    ))

    for sheet_name, df in sheets.items():
        if df.empty:
            continue

        # Truncate columns if too many
        if len(df.columns) > MAX_TABLE_COLS:
            df = df.iloc[:, :MAX_TABLE_COLS]

        headers = [_sanitize_cell(c) for c in df.columns]

        # Convert all data to strings
        all_rows = []
        for _, row in df.iterrows():
            all_rows.append([_sanitize_cell(v) for v in row])

        # Chunk rows into slides
        for chunk_idx in range(0, max(1, len(all_rows)), max_rows_per_slide):
            chunk = all_rows[chunk_idx : chunk_idx + max_rows_per_slide]
            page_num = chunk_idx // max_rows_per_slide + 1
            total_pages = (len(all_rows) + max_rows_per_slide - 1) // max_rows_per_slide

            if total_pages == 1:
                slide_title = sheet_name
            else:
                slide_title = f"{sheet_name} ({page_num}/{total_pages})"

            deck.slides.append(SlideContent(
                title=slide_title,
                slide_type="table",
                table_headers=headers,
                table_rows=chunk,
            ))

    return deck


def read_document(path: str) -> str:
    """Route to the right reader based on file extension."""
    ext = Path(path).suffix.lower()
    readers = {
        ".txt": read_txt,
        ".md": read_markdown,
        ".markdown": read_markdown,
        ".docx": read_docx,
        ".pdf": read_pdf,
        ".html": read_html,
        ".htm": read_html,
    }
    reader = readers.get(ext)
    if reader is None:
        raise ValueError(
            f"Unsupported file type: '{ext}'. "
            f"Supported: {', '.join(readers.keys())}"
        )
    try:
        src_bytes = os.path.getsize(path)
    except OSError:
        src_bytes = -1
    t0 = time.monotonic()
    text = reader(path)
    dt = time.monotonic() - t0
    n_lines = text.count("\n") + (0 if text.endswith("\n") or not text else 1)
    logger.info(
        "   Extracted %d chars (%d lines) from %s file (%d bytes) in %.2fs",
        len(text), n_lines, ext, src_bytes, dt,
    )
    logger.info("   Preview: %s", _preview(text, 240))
    logger.debug("=== EXTRACTED TEXT (%d chars) ===\n%s\n=== END EXTRACTED TEXT ===", len(text), text)
    return text
